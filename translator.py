from docx import Document
from deep_translator import GoogleTranslator,DeeplTranslator,ChatGptTranslator
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import deep_translator

doc_path = "input/IEC 62443-3-3 2013-image_table_content.docx"
font_name = 'B Nazanin'
font_size = 11
translators = {
    "google": lambda src,tgt: GoogleTranslator(source=src,target=tgt),
    "deepl": lambda src,tgt: DeeplTranslator(api_key="YOUR_KEY", target=tgt, source=src),
    "chatgpt": lambda src,tgt: ChatGptTranslator(api_key="YOUR_KEY",target=tgt,source=src),
}
output_format = "docx"  # or "pdf"


def set_run_rtl(run):
    r = run._r
    rPr = r.get_or_add_rPr()
    bidi = OxmlElement('w:rtl')
    #bidi.set(qn('w:rtl'))
    rPr.append(bidi)

def set_run_font(run,font_name,font_size):
    try:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.rFonts     
        rFonts.set(qn('w:cs'), font_name)
    except Exception:
        print("Error setting font for run")
        pass
    
# Set to correct english-persian paragraphs
def set_runs_rtl_and_font(paragraph):
    for run in paragraph.runs:
        set_run_rtl(run)
        set_run_font(run,font_name,font_size)



def set_paragraph_direction(paragraph, direction="LTR"):
    pPr = paragraph._p.get_or_add_pPr()

    # remove existing bidi element
    for el in pPr.xpath('./w:bidi'):
        pPr.remove(el)

    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), "1" if direction.upper() == "RTL" else "0")
    pPr.append(bidi)     


def translate_paragraphs():
    # Translate paragraphs
    for para in doc.paragraphs:
        print(f"Paragraph: {para.text}")
        translated = translator.translate(para.text)
        if translated:
            para.text = translated
            set_paragraph_direction(para, direction="RTL")       
            set_runs_rtl_and_font(para)       
                  
            
            

def translate_tables():    
    # Translate tables
    counter = 1
    for table in doc.tables:
        print("Table:")
        for row in table.rows:          
            for cell in row.cells:
                merge_count = cell._tc.get_or_add_tcPr().grid_span
                if counter < merge_count :# skip merged cells
                    counter += 1
                    continue

                if len(cell.text.strip()) <= 1:
                    continue

                counter = 1     
                translated = translator.translate(cell.text.strip())                
                print(f"Cell: {cell.text.strip()}")
                
                if translated:
                    cell.text = translated

                for paragraph in cell.paragraphs:                    
                    set_runs_rtl_and_font(paragraph)
                    set_paragraph_direction(paragraph, direction="RTL")
                   # paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT                    

        table.table_direction = 'rtl'
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER


def translate(doc_path):
    global doc
    doc = Document(doc_path)
    global translator
    translator = translators["google"](src='en', tgt='fa')
    translate_paragraphs()
    translate_tables()
    doc.save('output/translated.docx')

translate(doc_path)

