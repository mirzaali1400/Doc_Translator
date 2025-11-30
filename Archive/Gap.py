from openai import OpenAI
import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import datetime
import os


def gap(prompt):
    # Connect to your local Jan instance
    #client = OpenAI(base_url="https://api.gapgpt.app/v1", api_key="sk-TvPPujRByBU2ZItQ6KYIPD0Gq5xcQU8uZWcYvqDsSVBmaES2")
    client = OpenAI(base_url="http://10.116.83.20:1337/v1", api_key="test")
    response = client.chat.completions.create(
        #model="gpt-5", 
        model="Qwen3-14B-IQ4_XS",
        #model="Qwen3-42B-A3B-2507-Thinking-Abliterated-uncensored-TOTAL-RECALL-v2-Medium-MASTER-CODER_i1-IQ4_XS",
        messages=[
            {"role": "system", "content": "You are a translation assistant."},
            {"role": "user", "content": prompt},
        ]
    )

    res = response.choices[0].message.content

    return res


def translate(text):
    
    #prompt =  f'شما یک مترجم دقیق فارسی هستید. متن را به فارسی روزمره و محاوره‌ای که توسط فارسی‌زبانان استفاده می‌شود ترجمه کنید. از اصطلاحات طبیعی فارسی و لحن گفتگویی استفاده کنید و از ترجمه‌های رسمی یا تحت‌اللفظی خودداری کنید. ترجمه‌های شما باید طوری باشد که انگار از ابتدا به فارسی نوشته شده است. فقط ترجمه فارسی را خروجی دهید - بدون توضیحات، یادداشت‌ها یا متن اصلی.{text}'
    prompt = f"شما یک مترجم دقیق فارسی هستید. ترجمه متون ارائه شده را به صورت رسمی و به صورت روان انجام دهید{text}"
    response = gap(prompt)
    print(response)
    return response



def add_mixed_text(paragraph, text):
    # جداسازی کلمات/علائم انگلیسی و فارسی
    tokens = re.findall(r'[A-Za-z0-9\-_.]+|[^A-Za-z0-9\-_.]+', text)

    for tok in tokens:
        run = paragraph.add_run(tok)
        rPr = run._r.get_or_add_rPr()

        # اگر انگلیسی بود → LTR
        if re.match(r'[A-Za-z0-9]', tok):
            rtl = OxmlElement('w:rtl')
            rtl.set(qn('w:val'), '0')
            rPr.append(rtl)

        else:  # فارسی
            rtl = OxmlElement('w:rtl')
            rtl.set(qn('w:val'), '1')
            rPr.append(rtl)

        # فونت
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), 'B Nazanin')
        rFonts.set(qn('w:cs'), 'B Nazanin')
        rPr.append(rFonts)


def translatePDF_to_Word(file_name):
    doc = Document()

    time  = datetime.datetime.now().time()
    time = time.strftime("%H_%M_%S")

    # تنظیم فونت پاراگراف اصلی
    style = doc.styles['Normal']
    style.font.name = 'B Nazanin'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    style._element.rPr.rFonts.set(qn('w:cs'), 'B Nazanin')

    pdf = fitz.open(file_name)

    for page in pdf:
        blocks = sorted(page.get_text("blocks"), key=lambda b: (b[1], b[0]))

        for b in blocks:
            text = b[4].strip()
            if not text:
                continue

            text = text.replace("\n", " ")

            # ترجمه (تو اینجا تابع translate خودت را قرار بده)
            try:
                persian = translate(text)
            except Exception as e:
                print("Error:", e)
                continue

            # ---- ایجاد پاراگراف راست‌به‌چپ واقعی ----
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            pPr = p._p.get_or_add_pPr()

            bidi = OxmlElement('w:bidi')
            bidi.set(qn('w:val'), '1')
            pPr.append(bidi)

            rtl = OxmlElement('w:rtl')
            rtl.set(qn('w:val'), '1')
            pPr.append(rtl)

            # افزودن متن فارسی-انگلیسی بدون بهم ریختگی
            add_mixed_text(p, persian)

        doc.add_page_break()


    file_name = os.path.splitext(file_name)[0]
    doc.save(f"{file_name}_{time}_Translated.docx")
    print("فایل نهایی ذخیره شد →", file_name)





