from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import sys

# GitHub Copilot
# Requires: pip install python-docx

FONT_NAME = "B Nazain"   # user requested font name

def _ensure_rPr(run_elm):
    # run_elm is the CT_R element (run._element)
    try:
        return run_elm.get_or_add_rPr()
    except Exception:
        rPr = OxmlElement('w:rPr')
        run_elm.append(rPr)
        return rPr

def set_run_font_and_rtl(run, font_name=FONT_NAME):
    # high-level font name (for simple cases)
    run.font.name = font_name

    r = run._element
    rPr = _ensure_rPr(r)

    # set rFonts so Word uses the correct font for various script categories
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)

    # set run-level RTL
    rtl_run = OxmlElement('w:rtl')
    rtl_run.set(qn('w:val'), '1')
    rPr.append(rtl_run)

def set_paragraph_rtl(paragraph):
    # set paragraph alignment to right
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

    # set paragraph bidi property
    p = paragraph._p
    try:
        pPr = p.get_or_add_pPr()
    except Exception:
        pPr = OxmlElement('w:pPr')
        p.insert(0, pPr)
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def process_paragraphs(paragraphs):
    for para in paragraphs:
        set_paragraph_rtl(para)
        for run in para.runs:
            set_run_font_and_rtl(run)

def process_table(table):
    for row in table.rows:
        for cell in row.cells:
            process_paragraphs(cell.paragraphs)
            for inner_table in cell.tables:
                process_table(inner_table)

def process_document(doc):
    # body paragraphs
    process_paragraphs(doc.paragraphs)
    # tables in body
    for table in doc.tables:
        process_table(table)
    # headers/footers
    for section in doc.sections:
        header = section.header
        footer = section.footer
        process_paragraphs(header.paragraphs)
        for t in header.tables:
            process_table(t)
        process_paragraphs(footer.paragraphs)
        for t in footer.tables:
            process_table(t)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python setfont_run.py input.docx output.docx")
        sys.exit(1)
    in_path = sys.argv[1]
    out_path = sys.argv[2]

    doc = Document(in_path)
    process_document(doc)
    doc.save(out_path)
    print("Saved:", out_path)