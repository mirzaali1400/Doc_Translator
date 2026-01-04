from docx import Document
from deep_translator import GoogleTranslator,DeeplTranslator,ChatGptTranslator
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from colorama import Fore, Back, Style
from rich.progress import track
from tkinter import filedialog, messagebox
import os
from datetime import datetime
import time



#doc_path = "input/IEC 62443-3-3 2013-image_table_content.docx"
font_name = 'B Nazanin'
translator_name = "google"  # google, deepl, chatgpt
font_size = 11
OPENAI_API_KEY = os.environ["OPENAI_API_KEY"]
translators = {
    "google": lambda src,tgt: GoogleTranslator(source=src,target=tgt),
    "deepl": lambda src,tgt: DeeplTranslator(api_key="YOUR_KEY", target=tgt, source=src),
    "chatgpt": lambda src,tgt: ChatGptTranslator(api_key=OPENAI_API_KEY , target=tgt,source=src,model="gpt-5.1"),
}
output_format = "docx"  # or "pdf"


def set_run_rtl(run):
    r = run._r
    rPr = r.get_or_add_rPr()
    bidi = OxmlElement('w:rtl')  
    rPr.append(bidi)

def set_run_font(run,font_name,font_size):
    try:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.rFonts     
        rFonts.set(qn('w:cs'), font_name)
    except Exception as e:
        print(Fore.RED + Style.BRIGHT + "Error setting font for run")
        pass
    
# Set to correct english-persian paragraphs
def set_runs_rtl_and_font(paragraph):
    for run in paragraph.runs:
        set_run_rtl(run)
        set_run_font(run,font_name,font_size)



def set_paragraph_direction(paragraph, direction="LTR"):
    pPr = paragraph._p.get_or_add_pPr()

    # remove existing bidi element
    for el in pPr.xpath('./w:bidi'):
        pPr.remove(el)

    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), "1" if direction.upper() == "RTL" else "0")
    pPr.append(bidi)     


def translate_paragraphs():
    # Translate paragraphs
    para_count = len(doc.paragraphs)
    counter = 0
    for para in track(doc.paragraphs, description="[green]Translating paragraphs..."):  
        if para.text == "":
            continue
        if progress_callback:
            counter += 1
            progress_callback(counter / para_count * 50,"Translating paragraphs")      
        print(f"Paragraph: {para.text}")
        for i in range(1,3):
            try:
                translated = translator.translate(para.text)  
                break
            except:
                print(Fore.RED + Style.BRIGHT + "Translation error, retrying...")
                time.sleep(1)
                continue
        
        if translated:
            para.text = translated
            set_paragraph_direction(para, direction="RTL")       
            set_runs_rtl_and_font(para)       
                  
            
            

def translate_tables():    
    # Translate tables
    table_count = len(doc.tables)
    table_counter = 0
    counter = 1
    for table in track(doc.tables, description="[green]Translating tables..."):  
        if progress_callback:            
            progress_callback(50 + table_counter / table_count * 50,"Translating tables")      
            table_counter += 1          
        print("Table:")
        for row in track(table.rows, description="[green]Translating rows..."):          
            for cell in row.cells:
                merge_count = cell._tc.get_or_add_tcPr().grid_span
                if counter < merge_count :# skip merged cells
                    counter += 1
                    continue

                if len(cell.text.strip()) <= 1:
                    continue
                

                counter = 1 

                for i in range(1,3):
                    try:
                        translated = translator.translate(cell.text.strip())  
                        break
                    except:
                        print(Fore.RED + Style.BRIGHT + "Translation error, retrying...")
                        time.sleep(1)
                        continue  

                print(f"Cell: {cell.text.strip()}")
                
                if translated:
                    cell.text = translated

                for paragraph in cell.paragraphs:                    
                    set_runs_rtl_and_font(paragraph)
                    set_paragraph_direction(paragraph, direction="RTL")
                   # paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                          

        table.table_direction = 'rtl'
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER


def translate(doc_path,callback=None):

    global doc
    global progress_callback
    global translator

    doc = Document(doc_path)    
    translator = translators[translator_name](src='en', tgt='fa')    
    progress_callback = callback
    path,file_name = os.path.split(doc_path)
    file,ext = os.path.splitext(file_name)
    ouptput_path = os.path.join(path,f"{file}_translated_{translator_name}{ext}")
    
    translate_paragraphs()    
    translate_tables() 

    doc.save(ouptput_path)
    if progress_callback:
        progress_callback(100,"Translation Completed!")

if __name__ == "__main__":  

    doc_path = filedialog.askopenfilename(
        title="Select text file",
        filetypes=[("Word Files", "*.docx"), ("All Files", "*.*")]
    ) 

    if not doc_path:
        messagebox.showerror("Error", "No file selected.")
    else:
        tic = datetime.now()
        translate(doc_path)    
        toc = datetime.now()
        m,s = divmod((toc - tic).total_seconds(), 60)
        print(Fore.GREEN + Style.BRIGHT + f"Translation completed in {int(m)} : {int(s)} ")
