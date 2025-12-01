from docx import Document
from deep_translator import GoogleTranslator
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH



def set_run_rtl(run):
    r = run._r
    rPr = r.get_or_add_rPr()
    bidi = OxmlElement('w:rtl')
    #bidi.set(qn('w:rtl'))
    rPr.append(bidi)

def set_run_font(run,font_name='B Nazanin',font_size=10):
    try:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.rFonts     
        rFonts.set(qn('w:cs'), font_name)
    except Exception:
        print("Error setting font for run")
        pass
    

def set_runs_rtl_and_font(paragraph, font_name='B Nazanin', font_size=10):
    for run in paragraph.runs:
        set_run_rtl(run)
        set_run_font(run,font_name,font_size)



def set_paragraph_direction(paragraph, direction="LTR"):
    pPr = paragraph._p.get_or_add_pPr()

    # remove existing bidi element
    for el in pPr.xpath('./w:bidi'):
        pPr.remove(el)

    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), "1" if direction.upper() == "RTL" else "0")
    pPr.append(bidi)     
       


def translate(doc_path):

    print(f"Translating : {doc_path}")
    doc = Document(doc_path)
    translator = GoogleTranslator(source='en', target='fa')
    pre_cell_text = ""

    # Translate paragraphs
    for para in doc.paragraphs:
        print(f"Paragraph: {para.text}")
        translated = translator.translate(para.text)
        if translated:
            para.text = translated            
            set_paragraph_direction(para, direction="RTL")
            set_runs_rtl_and_font(para, font_name='B Nazanin', font_size=10)
            

    # Translate tables
    for table in doc.tables:
        print("Table:")
        for row in table.rows:          
            for cell in row.cells:
                if cell.text == pre_cell_text: 
                    continue
                pre_cell_text = cell.text
                translated = translator.translate(cell.text)                
                print(f"Cell: {cell.text}")
                

                if translated:
                    cell.text = translated

                for paragraph in cell.paragraphs:                    
                    for run in paragraph.runs:
                        set_run_rtl(run)
                        set_run_font(run,font_name='B Nazanin', font_size=10)

                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT                    

        table.table_direction = 'rtl'
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER
       
        
    
    doc.save('output/translated.docx')


def test():
    doc = Document(doc_path)
    i = 1
    for table in doc.tables:
        print("Table:")
        for row in table.rows:          
            print("count : ",len(row.cells))
            for cell in row.cells:
                
                #print("H cell vmerge : ",cell._tc.get_or_add_tcPr().grid_span)
                #print("V cell vmerge : ",cell._tc.get_or_add_tcPr().vMerge)          
                if  i <cell._tc.get_or_add_tcPr().grid_span :
                    i += 1
                    continue
                i = 1
                print(f"Cell: {cell.text}")
                    
                
                #for para in cell.paragraphs:
                  #  print(f"Cell Paragraph: {para.text}")



def test1():
    doc = Document(doc_path)
    for table in doc.tables:
        for row in table.rows:
            tr = row._tr
            for tc in tr.tc_lst:
                # -- vMerge="continue" indicates a spanned cell in a vertical merge --
                print("tc.vMerge: ", tc.vMerge)
                if tc.vMerge == "continue":
                    print("continue")
                    continue
                # --  --
                print("not continue")
#translate(doc_path)
test()

translate("input/IEC 62443-3-3 2013-image_table_content.docx")